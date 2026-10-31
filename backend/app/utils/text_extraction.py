import os
from pathlib import Path
from typing import Tuple, Dict, Any


def extract_text_from_pdf(file_path: str) -> Tuple[str, Dict[str, Any]]:
    """Extracts text content and metadata from a PDF file using PyMuPDF."""
    import pymupdf as fitz  # PyMuPDF
    text_content = []
    metadata = {}
    
    doc = fitz.open(file_path)
    metadata["page_count"] = len(doc)
    metadata["author"] = doc.metadata.get("author", "")
    metadata["title"] = doc.metadata.get("title", "")
    metadata["creationDate"] = doc.metadata.get("creationDate", "")
    
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text_content.append(page.get_text())
    doc.close()
    
    full_text = "\n".join(text_content).strip()
    return full_text, metadata


def extract_text_from_docx(file_path: str) -> Tuple[str, Dict[str, Any]]:
    """Extracts text content and metadata from a DOCX file."""
    import docx
    doc = docx.Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    metadata = {
        "paragraph_count": len(paragraphs),
        "table_count": len(doc.tables)
    }
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs.append(" | ".join(row_text))
                
    full_text = "\n".join(paragraphs).strip()
    return full_text, metadata


def extract_text_from_txt(file_path: str) -> Tuple[str, Dict[str, Any]]:
    """Extracts text from a plain text file."""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            content = f.read()
    metadata = {
        "character_count": len(content),
        "line_count": len(content.splitlines())
    }
    return content.strip(), metadata


def extract_file_content(file_path: str, file_type: str) -> Tuple[str, Dict[str, Any]]:
    """
    Dispatcher to extract text and metadata based on file type.
    Handles PDF, DOCX, TXT.
    For Image and Audio, returns clear extraction status without pretending unsupported OCR/Audio works.
    """
    file_type = file_type.lower()
    
    if file_type == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_type in ("docx", "doc"):
        return extract_text_from_docx(file_path)
    elif file_type in ("txt", "text", "csv", "log"):
        return extract_text_from_txt(file_path)
    elif file_type in ("png", "jpg", "jpeg", "webp"):
        # Image file
        file_size = os.path.getsize(file_path)
        metadata = {
            "media_type": "image",
            "file_size": file_size,
            "ocr_status": "OCR requires external OCR engine (e.g. Tesseract/Google Vision). Image preserved in secure storage."
        }
        return "[IMAGE EVIDENCE: Binary image file stored securely. Visual inspection or OCR required.]", metadata
    elif file_type in ("mp3", "wav", "m4a", "ogg", "aac"):
        # Audio file
        file_size = os.path.getsize(file_path)
        metadata = {
            "media_type": "audio",
            "file_size": file_size,
            "transcription_status": "Audio transcription requires configured speech-to-text pipeline. Audio file preserved in secure storage."
        }
        return "[AUDIO EVIDENCE: Audio recording stored securely. Manual or automated speech transcription required.]", metadata
    else:
        return "", {"note": f"Unsupported automated extraction for file type {file_type}"}
